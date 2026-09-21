from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from pathlib import Path

OUT = Path(__file__).resolve().parents[1] / 'docs' / 'AI_Plus_Project_Specification.docx'
OUT.parent.mkdir(exist_ok=True)
doc = Document()
sec = doc.sections[0]
sec.top_margin = sec.bottom_margin = Inches(.75)
sec.left_margin = sec.right_margin = Inches(.85)
styles = doc.styles
styles['Normal'].font.name = 'Aptos'; styles['Normal'].font.size = Pt(10.5)
for name, size in [('Title', 24), ('Heading 1', 16), ('Heading 2', 12)]:
    styles[name].font.name = 'Aptos Display'; styles[name].font.size = Pt(size); styles[name].font.color.rgb = RGBColor(0,0,0)

title = doc.add_paragraph('Tài liệu đặc tả dự án AI Plus', style='Title'); title.alignment = WD_ALIGN_PARAGRAPH.CENTER
# Remove the built-in Word Title style border so no blue rule appears in Word/LibreOffice.
pPr = title._p.get_or_add_pPr()
for border in pPr.findall(qn('w:pBdr')):
    pPr.remove(border)
pBdr = OxmlElement('w:pBdr'); bottom = OxmlElement('w:bottom'); bottom.set(qn('w:val'), 'nil'); pBdr.append(bottom); pPr.append(pBdr)
p = doc.add_paragraph('Tài liệu mô tả phạm vi, chức năng, bảo mật và lộ trình phát triển hệ thống AI Plus cho LSTS.'); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
doc.add_paragraph('Phiên bản: 1.0   |   Cập nhật: 18 September 2026').alignment = WD_ALIGN_PARAGRAPH.CENTER

def h(text, level=1): doc.add_heading(text, level=level)
def bullets(items):
    for x in items: doc.add_paragraph(x, style='List Bullet')
def table(headers, rows):
    t = doc.add_table(rows=1, cols=len(headers)); t.style='Table Grid'
    for i,x in enumerate(headers):
        c=t.rows[0].cells[i]; c.text=x
        tcPr=c._tc.get_or_add_tcPr(); sh=OxmlElement('w:shd'); sh.set(qn('w:fill'),'1F4E78'); tcPr.append(sh)
        for r in c.paragraphs[0].runs: r.font.color.rgb=RGBColor(255,255,255); r.font.bold=True
    for row in rows:
        cells=t.add_row().cells
        for i,x in enumerate(row): cells[i].text=x
    doc.add_paragraph('')

h('1. Tổng quan')
doc.add_paragraph('AI Plus là nền tảng AI nội bộ của LSTS, hỗ trợ giáo viên và nhân sự trao đổi với AI, làm việc cùng Agent chuyên biệt, phân tích tài liệu và tạo đầu ra phục vụ công việc.')
h('2. Mục tiêu')
bullets(['Cung cấp AI an toàn trong môi trường trường học.', 'Hỗ trợ phân tích tài liệu, soạn nội dung và tạo file.', 'Kiểm soát quyền truy cập, quota token và dữ liệu nhạy cảm.', 'Mở rộng dần sang các quy trình công việc có xác nhận của user.'])
h('3. Người dùng và quyền')
table(['Nhóm', 'Quyền chính'], [('User / Staff', 'Chat, Agent cá nhân, knowledge, showcase, quota và artifact cá nhân.'), ('Admin', 'Quản lý user, prompt, template, showcase, ticket và audit log.')])
h('4. Chức năng chính')
h('AI Chat và Agent Workspace', 2)
bullets(['Chat riêng hoặc chat với Agent, lưu lịch sử và streaming phản hồi.', 'Upload ảnh, Word, Excel, PDF, CSV và tài liệu hỗ trợ.', 'Sidebar chat sắp xếp theo cập nhật gần nhất.', 'Knowledge Agent sử dụng keyword RAG; PDF scan được render thành ảnh cho AI vision.'])
h('Sharing Showcase và Admin', 2)
bullets(['Agent chia sẻ xuất hiện tại Sharing and Showcase; user khác có thể dùng để tạo bản riêng.', 'Admin Panel quản lý user, prompt, template, showcase, FAQ, support ticket và audit log.'])
h('5. Giai đoạn 1 Tạo file và Email nháp')
table(['Khả năng', 'Kết quả'], [('Tạo Excel', 'File .xlsx private để user tải xuống.'), ('Tạo Word', 'File .docx private để user tải xuống.'), ('Tạo PDF', 'File .pdf private để user tải xuống.'), ('Email nháp', 'Subject và body được lưu, không gửi email thật.')])
doc.add_paragraph('Luồng: user trao đổi với AI hoặc upload tài liệu → yêu cầu tạo file hoặc email nháp → AI tạo nội dung → backend tạo artifact private → chat hiển thị link tải → hệ thống lưu audit log.')
h('6. Bảo mật và quyền riêng tư')
bullets(['Kiểm tra ownership cho conversation, attachment và artifact.', 'Artifact và attachment dùng storage private, không public trực tiếp.', 'Rate limit cho login, chat, upload Agent và support request.', 'Sanitize Markdown AI và escape nội dung admin để giảm rủi ro XSS.', 'Lọc PII: email ngoài @lsts.edu.vn, số điện thoại, mã học sinh, CCCD/CMND, địa chỉ cụ thể, số tài khoản, hộ chiếu và biển số xe.'])
h('7. Usage quota')
table(['Giai đoạn', 'Quota'], [('Testing', '20 triệu tokens mỗi user mỗi tháng'), ('Training', '10 triệu tokens mỗi user mỗi tháng')])
doc.add_paragraph('Quota được reset vào ngày đầu tháng. Conversation đã xóa không hiển thị trong Recent Activity.')
h('8. Roadmap')
table(['Giai đoạn', 'Phạm vi'], [('Giai đoạn 1', 'Tạo Excel/Word/PDF, email nháp, artifact download, streaming và audit log.'), ('Giai đoạn 2', 'Microsoft 365/OneDrive, Outlook draft, queue, object storage và virus scanning.'), ('Giai đoạn 3', 'Gửi email thật sau xác nhận, workflow nhiều bước, template chính thức và phân quyền tool.')])
h('9. Yêu cầu deploy')
bullets(['HTTPS và reverse proxy hỗ trợ SSE streaming.', 'Queue worker, scheduler, database backup và private file storage.', 'Poppler pdftoppm cho PDF scan.', 'Cấu hình upload limit, .env, cache và monitoring log.'])
doc.save(OUT)
print(OUT)
