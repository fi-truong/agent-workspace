from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from pathlib import Path

out = Path(__file__).resolve().parents[1] / 'docs' / 'AI_Plus_Project_Specification_EN.docx'
d = Document(); s=d.sections[0]; s.top_margin=s.bottom_margin=Inches(.75); s.left_margin=s.right_margin=Inches(.85)
d.styles['Normal'].font.name='Aptos'; d.styles['Normal'].font.size=Pt(10.5)
for n,size in [('Title',24),('Heading 1',16),('Heading 2',12)]: d.styles[n].font.name='Aptos Display'; d.styles[n].font.size=Pt(size); d.styles[n].font.color.rgb=RGBColor(0,0,0)
p=d.add_paragraph('AI Plus Project Specification',style='Title'); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
pp=p._p.get_or_add_pPr(); bd=OxmlElement('w:pBdr'); b=OxmlElement('w:bottom'); b.set(qn('w:val'),'nil'); bd.append(b); pp.append(bd)
for text in ['This document defines the scope, capabilities, security model and development roadmap for the LSTS AI Plus platform.','Version 1.0   |   Updated 18 September 2026']:
 q=d.add_paragraph(text); q.alignment=WD_ALIGN_PARAGRAPH.CENTER
def h(x,l=1): d.add_heading(x,l)
def bs(xs):
 for x in xs:d.add_paragraph(x,style='List Bullet')
def t(headers,rows):
 x=d.add_table(rows=1,cols=len(headers));x.style='Table Grid'
 for i,v in enumerate(headers):
  c=x.rows[0].cells[i];c.text=v; z=c._tc.get_or_add_tcPr();sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'1F4E78');z.append(sh)
 for r in c.paragraphs[0].runs:r.font.color.rgb=RGBColor(255,255,255);r.font.bold=True
 for row in rows:
  cells=x.add_row().cells
  for i,v in enumerate(row):cells[i].text=v
 d.add_paragraph()
h('1. Overview');d.add_paragraph('AI Plus is LSTS’s internal AI platform for teachers and staff. It combines AI chat, specialised Agents, document analysis, central administration and work-output generation.')
h('2. Objectives');bs(['Provide secure AI assistance in a school environment.','Support document analysis, content drafting and file creation.','Control access rights, token quotas and sensitive data.','Expand from chat into user-confirmed work processes.'])
h('3. Users and permissions');t(['Group','Primary permissions'],[('User / Staff','Chat, personal Agents, knowledge files, showcase, quota and personal artifacts.'),('Admin','Manage users, prompts, templates, showcase, tickets and audit logs.')])
h('4. Core capabilities');h('AI Chat and Agent Workspace',2);bs(['Standalone or Agent chat with persisted history and streaming responses.','Image, Word, Excel, PDF, CSV and supported document upload.','Conversations ordered by most recent update.','Keyword RAG for Agent knowledge; scanned PDFs are rendered as images for AI vision.'])
h('Sharing Showcase and Administration',2);bs(['Shared Agents appear in Sharing and Showcase and can be copied for personal use.','The Admin Panel manages users, prompts, templates, showcase content, FAQs, tickets and audit logs.'])
h('5. Phase 1 File Creation and Email Drafts');t(['Capability','Outcome'],[('Create Excel','Private .xlsx file available to its owner.'),('Create Word','Private .docx file available to its owner.'),('Create PDF','Private .pdf file available to its owner.'),('Email draft','Subject and body are saved; no email is sent.')]);d.add_paragraph('Flow: user chats with AI or uploads documents → requests a file or email draft → AI creates the content → backend creates a private artifact → chat shows a download link → the system records an audit event.')
h('6. Security and privacy');bs(['Ownership checks for conversations, attachments and artifacts.','Private storage for attachments and artifacts; no direct public access.','Rate limits for login, chat, Agent uploads and support requests.','Sanitised AI Markdown and escaped admin content to reduce XSS risk.','PII filtering for external email, phone numbers, student IDs, national IDs, specific addresses, bank accounts, passports and licence plates.'])
h('7. Usage quota');t(['Phase','Quota'],[('Testing','20 million tokens per user per month'),('Training','10 million tokens per user per month')]);d.add_paragraph('Quotas reset on the first day of each month. Deleted conversations are hidden from Recent Activity.')
h('8. Roadmap');t(['Phase','Scope'],[('Phase 1','Excel, Word and PDF creation, email drafts, artifact downloads, streaming and audit logs.'),('Phase 2','Microsoft 365 and OneDrive, Outlook drafts, queues, object storage and virus scanning.'),('Phase 3','Confirmed email sending, multi-step workflows, official templates and tool permissions.')])
h('9. Deployment requirements');bs(['HTTPS and a reverse proxy that supports SSE streaming.','Queue worker, scheduler, database backups and private file storage.','Poppler pdftoppm for scanned PDF support.','Appropriate upload limits, environment configuration, caching and log monitoring.'])
d.save(out);print(out)
