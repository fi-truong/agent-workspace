from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from pathlib import Path
import re

ROOT=Path(__file__).resolve().parents[1]; OUT=ROOT/'docs'; OUT.mkdir(exist_ok=True)
def style(doc,title,subtitle):
 s=doc.sections[0];s.top_margin=s.bottom_margin=Inches(.75);s.left_margin=s.right_margin=Inches(.85)
 doc.styles['Normal'].font.name='Aptos';doc.styles['Normal'].font.size=Pt(10)
 for n,z in [('Title',24),('Heading 1',16),('Heading 2',12)]:doc.styles[n].font.name='Aptos Display';doc.styles[n].font.size=Pt(z);doc.styles[n].font.color.rgb=RGBColor(0,0,0)
 p=doc.add_paragraph(title,style='Title');p.alignment=WD_ALIGN_PARAGRAPH.CENTER;ppr=p._p.get_or_add_pPr();bd=OxmlElement('w:pBdr');bt=OxmlElement('w:bottom');bt.set(qn('w:val'),'nil');bd.append(bt);ppr.append(bd)
 for x in [subtitle,'Version 1.0   |   Updated 18 September 2026']:
  q=doc.add_paragraph(x);q.alignment=WD_ALIGN_PARAGRAPH.CENTER
def add_table(doc, lines):
 rows=[[c.strip() for c in x.strip().strip('|').split('|')] for x in lines if not re.match(r'^\|\s*-+',x)]
 t=doc.add_table(rows=1,cols=len(rows[0]));t.style='Table Grid'
 for i,v in enumerate(rows[0]):
  c=t.rows[0].cells[i];c.text=v;pr=c._tc.get_or_add_tcPr();sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'1F4E78');pr.append(sh)
  for r in c.paragraphs[0].runs:r.font.bold=True;r.font.color.rgb=RGBColor(255,255,255)
 for row in rows[1:]:
  cells=t.add_row().cells
  for i,v in enumerate(row):cells[i].text=v
 doc.add_paragraph()
def markdown_doc(path,out,title,subtitle):
 d=Document();style(d,title,subtitle);lines=path.read_text().splitlines();i=0
 while i<len(lines):
  x=lines[i].strip()
  if not x:i+=1;continue
  if x.startswith('# '):i+=1;continue
  if x.startswith('## '):d.add_heading(x[3:],1)
  elif x.startswith('### '):d.add_heading(x[4:],2)
  elif x.startswith('- '):d.add_paragraph(x[2:],style='List Bullet')
  elif x.startswith('|'):
   a=[]
   while i<len(lines) and lines[i].strip().startswith('|'):a.append(lines[i].strip());i+=1
   add_table(d,a);continue
  else:d.add_paragraph(x)
  i+=1
 d.save(out)

vi=ROOT/'PROJECT_SPEC.md';markdown_doc(vi,OUT/'AI_Plus_Project_Specification_Full_VI.docx','Tài liệu đặc tả dự án AI Plus','Bản đầy đủ: phạm vi, chức năng, bảo mật, dữ liệu, API và lộ trình')

en='''
1. Overview
AI Plus is LSTS’s internal AI platform for teachers and staff. It supports academic, operational and administrative work through AI chat, specialised Agents, knowledge retrieval and controlled work outputs.
2. Objectives
• Provide secure AI assistance in a school environment.
• Enable users to discuss, analyse documents and create content.
• Enable purpose-built Agents and school-wide sharing.
• Control access rights, token quotas and sensitive data.
• Progress from chat to AI-supported work processes.
3. Users and Permissions
User and Staff: chat, personal Agents, knowledge upload, showcase, quota and personal artifacts.
Admin: manage users, prompts, templates, showcase, support tickets and audit logs.
All conversations, attachments, artifacts and email drafts belong to their creator. Other users cannot access them without an explicit sharing mechanism.
4. Existing Modules
AI Chat: standalone or Agent chat; saved conversation history; real-time streaming; automatic conversation title; sidebar ordered by latest update; image, Word, Excel, PDF, CSV and supported-document upload. Text PDFs are extracted directly. Scanned or image-only PDFs are rendered as images for AI vision.
Agent Workspace and RAG: create, edit and delete Agents; configure title, description, system prompt and knowledge files; attach an Agent to a conversation; keyword RAG retrieval from Agent knowledge; share Agents school-wide through Sharing and Showcase.
Sharing and Showcase: Agents selected as Share with school appear in the showcase. Other users can use an Agent to create a personal copy. New marks recent content; Popular marks heavily used Agents.
Admin Panel: manage users, roles, departments and account status; prompts, templates, showcase, FAQs and support tickets; audit logs and an Admin shortcut from AI Plus.
Usage quota: prompt and token usage; monthly reset; Testing allowance is 20 million tokens per user per month; Training allowance is 10 million; deleted conversations are hidden from Recent Activity.
5. AI Tools and Artifacts Phase 1
Users can request private Excel (.xlsx), Word (.docx) and PDF (.pdf) artifacts from the current AI content. Flow: user chats or uploads documents; requests a file or email draft; AI produces content; backend creates a private artifact linked to the user and conversation; chat shows a download link; artifact appears in recent file history. Original user files are never overwritten.
Email drafts store an AI-generated subject and body. Phase 1 does not send email, connect mailboxes or send externally. SSE reports AI response and file/draft generation. Artifact creation and email draft creation are audit logged. Artifact download verifies ownership.
6. Main Data Model
users: accounts, roles, departments and active state.
conversations: conversations, optionally linked to an Agent.
messages: user and assistant messages with token usage.
agents and knowledge_chunks: Agent configuration, knowledge and RAG chunks.
usage_logs: token and activity tracking.
ai_artifacts: generated file owner, conversation, path, MIME type and size.
email_drafts: draft owner, conversation, recipients, subject, body and attachment metadata.
admin_audit_logs: administrative and AI tool events.
7. Key Routes
POST /ai-plus/agent-workspace/send: non-streaming chat.
POST /ai-plus/agent-workspace/send-stream: SSE streaming chat.
GET /ai-plus/artifacts/{artifact}/download: owner-authorised artifact download.
GET /ai-plus/agent-workspace/attachments/{conversation}/{filename}: owner-authorised attachment download.
8. Security and Privacy
Authentication is required for functional AI routes. Ownership is checked for conversations, attachments and artifacts. Attachments and artifacts use private storage. Login, chat, Agent uploads and support requests are rate limited. AI Markdown is sanitised and dynamic admin content is escaped to reduce XSS risk. The model cannot run arbitrary system commands; every action must pass through a defined backend tool.
PII filtering applies before content is sent to AI: email outside @lsts.edu.vn, Vietnamese phone numbers, student IDs, national IDs, specific addresses, bank accounts, passports and licence plates. Exact @lsts.edu.vn email is allowed for internal work. Raw PII is not logged in audit events.
9. Current Limits
The system does not yet send email, connect Microsoft 365, OneDrive, Google Drive or Gmail, or edit a user’s original file directly. Artifacts are generated from AI responses; advanced official Word and Excel templates are a later step. Scanned PDF page counts are limited by configuration. Large file jobs should move to queues in production.
10. Roadmap
Phase 1: Excel, Word and PDF creation; email drafts; private artifact downloads, history, streaming and audit logs; scanned PDF fallback.
Phase 2: Microsoft 365 and OneDrive integration, saving artifacts to OneDrive, Outlook drafts, user confirmation before sending, queue workers, Redis/object storage and virus scanning.
Phase 3: confirmed real email sending, multi-step workflows, official Word and Excel templates, tool permissions by role and department, and advanced usage and audit dashboards.
11. Deployment Requirements
HTTPS and a reverse proxy configured for SSE streaming; Laravel queue workers and scheduler; production database backup and private file storage; Poppler pdftoppm for scanned PDF support; PHP/Nginx upload limits; environment configuration, cache optimisation and log/error monitoring.
'''.strip().splitlines()
d=Document();style(d,'AI Plus Project Specification','Full specification: scope, capabilities, security, data model, routes and roadmap')
for line in en:
 line=line.strip()
 if not line:continue
 if re.match(r'^\d+\. ',line):d.add_heading(line,1)
 elif line.startswith('• '):d.add_paragraph(line[2:],style='List Bullet')
 else:d.add_paragraph(line)
# Use the full English specification, mirroring the Vietnamese document's headings, tables and detail.
markdown_doc(ROOT/'PROJECT_SPEC_EN.md', OUT/'AI_Plus_Project_Specification_Full_EN.docx', 'AI Plus Project Specification', 'Full specification: scope, capabilities, security, data model, routes and roadmap')
